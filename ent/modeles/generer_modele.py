"""Génère le modèle Excel d'import de l'ENT (onglets 1 à 10), à remplir par l'école puis à importer
dans Administration → Importer des données. Aucune donnée réelle : seulement les en-têtes, une ligne
d'exemple fictive (commençant par « Exemple », ignorée à l'import) et des listes déroulantes.
Usage : python3 generer_modele.py  →  modele-import-chateaubriand.xlsx"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

NAVY = '003366'
CLASSES = ['6e', '5e', '4e', '3e', '2de', '1re Générale', '1re STMG', 'Terminale Générale', 'Terminale STMG']
LISTS = {
    'classe': CLASSES,
    'sexe': ['F', 'M'],
    'nationalite': ['Malgache', 'Française', 'Chinoise', 'Malgache et française', 'Autre'],
    'lien': ['Mère', 'Père', 'Tuteur / tutrice', 'Autre'],
    'ouinon': ['Oui', 'Non'],
    'fonction': ['Professeur', 'CPE', 'Secrétariat', 'Direction', 'Autre'],
    'jour': ['Lundi', 'Mardi', 'Mercredi', 'Jeudi', 'Vendredi', 'Samedi'],
    'mode': ['Espèces', 'Chèque', 'Virement', 'Mobile money'],
    'categorie': ['Mobilier', 'Équipement', 'Fournitures', 'Trousse de secours'],
    'etat': ['Neuf', 'Bon état', 'Usé', 'Déchiré / abîmé'],
    'typemanuel': ['Manuel (à rendre)', 'Cahier (gardé par l\'élève)'],
}
# (titre de l'onglet, explication, colonnes : (en-tête, largeur, liste, commentaire), ligne d'exemple)
SHEETS = [
    ('1 Élèves', 'Une ligne par élève inscrit cette année.', [
        ('N° (matricule)', 14, None, 'Ex. TG-001. Laisser vide : l\'ENT propose un numéro.'),
        ('Nom', 18, None, None), ('Prénom(s)', 22, None, 'Tous les prénoms, comme sur l\'acte de naissance.'),
        ('Prénom d\'usage', 16, None, 'Le prénom utilisé au quotidien, s\'il y en a plusieurs.'),
        ('Sexe', 7, 'sexe', None), ('Née le', 13, None, 'JJ/MM/AAAA'), ('Lieu de naissance', 18, None, None),
        ('Nationalité', 16, 'nationalite', 'Autre : préciser dans la colonne suivante.'), ('Autre nationalité', 16, None, None),
        ('Classe', 18, 'classe', None), ('Spécialités', 22, None, '1re/Tle générales : ex. Maths, PC, SVT'),
        ('Date d\'entrée', 13, None, 'JJ/MM/AAAA ou année d\'entrée'), ('Établissement précédent', 22, None, None),
        ('Téléphone (élève)', 16, None, None), ('Adresse e-mail', 24, None, None),
        ('Allergies', 18, None, 'Dont allergie au paracétamol.'), ('Traitement / maladie', 20, None, None)],
        ['Exemple', 'RAKOTO', 'Hery Jean', 'Hery', 'M', '14/03/2009', 'Toamasina', 'Malgache', '', 'Terminale STMG', '', '2021', 'EPP Tamatave', '034 00 000 00', 'hery@exemple.mg', 'Aucune', '']),
    ('2 Responsables', 'Une ligne par responsable (deux lignes si deux parents). Chaque parent aura son propre accès.', [
        ('N° élève (matricule)', 16, None, 'Le même numéro que dans l\'onglet 1.'), ('Nom de l\'élève', 22, None, None),
        ('Lien', 14, 'lien', None), ('Nom', 18, None, None), ('Prénom', 18, None, None),
        ('Téléphone', 16, None, None), ('Téléphone 2', 16, None, None), ('E-mail', 24, None, 'Sert à recevoir les codes et les messages.'),
        ('Profession', 18, None, None), ('À appeler en urgence', 12, 'ouinon', None), ('Peut venir chercher l\'élève', 14, 'ouinon', None)],
        ['Exemple', 'RAKOTO Hery', 'Mère', 'RAKOTO', 'Voahangy', '032 00 000 00', '', 'parent@exemple.mg', 'Commerçante', 'Oui', 'Oui']),
    ('3 Personnel', 'Professeurs et personnel. Pas les taux horaires : ils se saisissent dans Salaires.', [
        ('Nom', 18, None, None), ('Prénom', 18, None, None), ('Fonction', 14, 'fonction', None),
        ('Matière(s)', 26, None, 'Séparées par des virgules. Ex. Management, Droit'),
        ('Classe(s)', 30, None, 'Séparées par des virgules. Ex. 1re STMG, Terminale STMG'),
        ('Téléphone', 16, None, None), ('E-mail', 24, None, None)],
        ['Exemple', 'Lova', 'Professeur', 'Management', '1re STMG, Terminale STMG', '034 00 000 00', 'prof@exemple.mg']),
    ('4 Emploi du temps', 'Une ligne par cours de la semaine type.', [
        ('Classe', 18, 'classe', None), ('Jour', 11, 'jour', None), ('Début', 8, None, 'HH:MM'), ('Fin', 8, None, 'HH:MM'),
        ('Matière', 22, None, None), ('Professeur', 22, None, None), ('Salle', 12, None, None)],
        ['Exemple', 'Lundi', '07:00', '09:00', 'Management', 'Mme Exemple', 'Salle 12']),
    ('5 Notes', 'Notes de l\'année en cours, une ligne par élève et par évaluation.', [
        ('N° élève (matricule)', 16, None, None), ('Nom de l\'élève', 22, None, None), ('Classe', 18, 'classe', None),
        ('Matière', 20, None, None), ('Date', 12, None, 'JJ/MM/AAAA'), ('Intitulé', 24, None, 'Ex. DS n°1'),
        ('Note', 8, None, 'Nombre, ou « abs »'), ('Sur', 6, None, '20 par défaut'), ('Coefficient', 10, None, None)],
        ['Exemple', 'RAKOTO Hery', 'Terminale STMG', 'Management', '24/09/2026', 'DS n°1', '14,5', '20', '2']),
    ('6 Notes officielles 1re', 'Élèves de Terminale : notes OFFICIELLES obtenues en fin de 1re (relevé de juin). Sert à la projection au bac.', [
        ('N° élève (matricule)', 16, None, None), ('Nom de l\'élève', 22, None, None),
        ('Français écrit', 10, None, None), ('Français oral', 10, None, None), ('Mathématiques (anticipée)', 13, None, None),
        ('Spécialité abandonnée (ou SGN)', 15, None, None), ('Histoire-géo et EMC', 12, None, None),
        ('Ens. scientifique (ou maths STMG)', 15, None, None), ('LVA', 8, None, None), ('LVB', 8, None, None)],
        ['Exemple', 'RAKOTO Hery', '11', '13', '12', '10,5', '12', '9', '14', '12']),
    ('7 Paiements', 'Paiements déjà reçus cette année.', [
        ('N° élève (matricule)', 16, None, None), ('Nom de l\'élève', 22, None, None), ('Article', 24, None, 'Ex. Frais de scolarité octobre'),
        ('Montant (Ar)', 13, None, None), ('Date', 12, None, 'JJ/MM/AAAA'), ('Mode', 13, 'mode', None), ('N° de reçu', 12, None, None), ('Payé par', 18, None, None)],
        ['Exemple', 'RAKOTO Hery', 'Frais de scolarité septembre', '450000', '05/09/2026', 'Espèces', '000123', 'Mère']),
    ('8 Manuels', 'Stock de manuels et de cahiers, par état.', [
        ('Titre', 30, None, None), ('Matière', 18, None, None), ('Niveau(x)', 22, None, 'Séparés par des virgules'),
        ('Type', 22, 'typemanuel', None), ('Neuf', 7, None, None), ('Bon état', 9, None, None), ('Usé', 7, None, None), ('Déchiré / abîmé', 10, None, None),
        ('Prix de location ou de vente (Ar)', 16, None, None), ('Prix en cas de perte (Ar)', 14, None, None)],
        ['Exemple', 'Management', 'Terminale STMG', 'Manuel (à rendre)', '2', '5', '1', '0', '25000', '72000']),
    ('9 Prêts de manuels', 'Manuels remis aux élèves cette année.', [
        ('N° élève (matricule)', 16, None, None), ('Nom de l\'élève', 22, None, None), ('Titre du manuel', 30, None, None),
        ('État au départ', 14, 'etat', None), ('Date de remise', 13, None, 'JJ/MM/AAAA')],
        ['Exemple', 'RAKOTO Hery', 'Management Terminale STMG', 'Bon état', '07/09/2026']),
    ('10 Inventaire', 'Mobilier, équipement, fournitures et trousse de secours.', [
        ('Catégorie', 18, 'categorie', None), ('Article', 28, None, None), ('Quantité', 9, None, None), ('Dont cassés', 10, None, None),
        ('Lieu', 18, None, None), ('Seuil de recommande', 12, None, 'En dessous : à recommander'), ('Péremption', 12, None, 'Trousse de secours : JJ/MM/AAAA')],
        ['Exemple', 'Table élève double', '40', '1', 'Salles de classe', '2', '']),
]

wb = Workbook()
info = wb.active
info.title = 'Lisez-moi'
info['A1'] = 'Modèle d\'import — ENT École Collège Lycée Chateaubriand'
info['A1'].font = Font(bold=True, size=14, color=NAVY)
lines = [
    'Remplissez les onglets dans l\'ordre ; seuls les onglets 1 à 4 sont nécessaires pour démarrer.',
    'Ne changez pas les titres des colonnes. Une ligne = un élève, un parent, un cours…',
    'La ligne grise « Exemple » montre le format attendu : elle est ignorée à l\'import (vous pouvez l\'effacer).',
    'Les colonnes avec une flèche proposent une liste : choisissez dans la liste.',
    'Dates : JJ/MM/AAAA. Montants : en Ariary, sans espace ni « Ar ».',
    'Import : Administration → Importer des données → « Importer le modèle Excel ».',
    'Confidentialité : ce fichier contient des données d\'élèves mineurs. Ne l\'envoyez à personne ;',
    'gardez-le sur l\'ordinateur de l\'école et importez-le vous-même.',
]
for i, t in enumerate(lines, start=3):
    info.cell(row=i, column=1, value=t)
info.cell(row=12, column=1, value='Onglets').font = Font(bold=True)
for i, (title, expl, _, _) in enumerate(SHEETS, start=13):
    info.cell(row=i, column=1, value=f'{title} — {expl}')
info.column_dimensions['A'].width = 120

head_font = Font(bold=True, color='FFFFFF')
head_fill = PatternFill('solid', fgColor=NAVY)
ex_font = Font(italic=True, color='7F8C9A')
thin = Side(style='thin', color='C8D3E0')
for title, expl, cols, example in SHEETS:
    ws = wb.create_sheet(title)
    for c, (name, width, lst, note) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        cell.border = Border(bottom=thin)
        if note:
            cell.comment = Comment(note, 'ENT')
        ws.column_dimensions[cell.column_letter].width = width
        if lst:
            dv = DataValidation(type='list', formula1='"' + ','.join(LISTS[lst]) + '"', allow_blank=True)
            dv.error = 'Choisissez une valeur dans la liste.'
            ws.add_data_validation(dv)
            dv.add(f'{cell.column_letter}2:{cell.column_letter}2000')
    for c, v in enumerate(example, start=1):
        ws.cell(row=2, column=c, value=v).font = ex_font
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'A2'

wb.save('modele-import-chateaubriand.xlsx')
print('modele-import-chateaubriand.xlsx')
